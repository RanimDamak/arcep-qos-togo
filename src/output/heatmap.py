"""Heatmap - Etape 9 du document technique (SSVI). Matrice n x 24 (n forfaits x 24 heures),
construite depuis evenement_detecte (type_evenement='ANOMALIE'). Cellule coloree si anomalie :
    - Rouge  : niveau=2 (depassement 3-sigma)
    - Orange : niveau=1 (depassement 2-sigma, surveillance)
    - Blanc  : pas d'anomalie (ou pas de profil suffisant pour juger, cf. DetecteurAnomalie)

Seuls les forfaits ayant eu AU MOINS UNE ligne trafic_horaire sur le jour sont inclus dans la matrice (pas tout
le catalogue `forfait` - un forfait jamais actif ce jour-la n'a pas sa place dans une heatmap journaliere).
"""
from __future__ import annotations

import argparse
import datetime as dt
from typing import Optional

COULEUR_NIVEAU_2 = "FFC7CE"  # rouge clair
COULEUR_NIVEAU_1 = "FFEB9C"  # orange clair
COULEUR_NORMAL   = "FFFFFF"  # blanc


class Heatmap:
    """Construit et exporte la heatmap n x 24 des anomalies detectees pour un jour donne."""

    def construire(self, jour: dt.date) -> dict:
        """Retourne {forfait_id: {heure: niveau_ou_None}} pour tous les forfaits actifs ce jour-la."""
        from src.donnees.db import connexion

        debut = dt.datetime.combine(jour, dt.time.min, tzinfo=dt.timezone.utc)
        fin = debut + dt.timedelta(days=1)

        with connexion() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT DISTINCT forfait_id FROM trafic_horaire WHERE bucket >= %s AND bucket < %s ORDER BY forfait_id",
                    (debut, fin),
                )
                forfaits = [r[0] for r in cur.fetchall()]

                cur.execute(
                    """
                    SELECT forfait_id, EXTRACT(HOUR FROM bucket)::int, niveau
                    FROM evenement_detecte
                    WHERE type_evenement = 'ANOMALIE' AND bucket >= %s AND bucket < %s
                    """,
                    (debut, fin),
                )
                anomalies = {(r[0], r[1]): r[2] for r in cur.fetchall()}

        matrice = {
            fid: {h: anomalies.get((fid, h)) for h in range(24)}
            for fid in forfaits
        }
        return matrice

    def exporter_excel(self, matrice: dict, chemin: str) -> None:
        """Ecrit la heatmap dans un fichier .xlsx, forfaits en lignes, heures 0-23 en colonnes."""
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Heatmap"

        ws.cell(row=1, column=1, value="forfait_id")
        for h in range(24):
            ws.cell(row=1, column=h + 2, value=h)

        fills = {
            2: PatternFill(start_color=COULEUR_NIVEAU_2, end_color=COULEUR_NIVEAU_2, fill_type="solid"),
            1: PatternFill(start_color=COULEUR_NIVEAU_1, end_color=COULEUR_NIVEAU_1, fill_type="solid"),
            None: PatternFill(start_color=COULEUR_NORMAL, end_color=COULEUR_NORMAL, fill_type="solid"),
        }

        for i, (forfait_id, heures) in enumerate(sorted(matrice.items()), start=2):
            ws.cell(row=i, column=1, value=forfait_id)
            for h in range(24):
                niveau = heures.get(h)
                cell = ws.cell(row=i, column=h + 2, value=niveau if niveau is not None else "")
                cell.fill = fills[niveau]

        wb.save(chemin)


def _parse_date(valeur: str) -> dt.date:
    return dt.date.fromisoformat(valeur)


def main(argv: Optional[list[str]] = None) -> None:
    p = argparse.ArgumentParser(description="Genere la heatmap des anomalies (n forfaits x 24h) pour un jour donne.")
    p.add_argument("--jour", required=True, help="Jour a analyser, ISO (ex. 2025-08-02)")
    p.add_argument("--sortie", default="heatmap.xlsx", help="Chemin du fichier .xlsx de sortie")
    args = p.parse_args(argv)

    heatmap = Heatmap()
    matrice = heatmap.construire(_parse_date(args.jour))
    heatmap.exporter_excel(matrice, args.sortie)

    n_anomalies = sum(1 for h in matrice.values() for v in h.values() if v is not None)
    print(f"Heatmap generee : {len(matrice)} forfait(s) x 24h -> {args.sortie} ({n_anomalies} cellule(s) anormale(s))")


if __name__ == "__main__":
    main()