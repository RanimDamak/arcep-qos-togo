"""Rapport - Etape 9 (partie 2). Genere un rapport Excel par operateur : heatmap, detail des anomalies, alertes QoS, correlations. Log l'execution dans la table `rapport`.

Important : les alertes QoS (congestion/QoS/half-rate) sont RESEAU ENTIER (cf. congestion_horaire, QUALCOP ne distingue pas les operateurs) - cette feuille
n'est donc PAS filtree par operateur, contrairement aux anomalies et correlations (qui, elles, sont attachees a un forfait_id, donc a un operateur via
forfait.operateur_id). Le rapport le precise explicitement pour eviter toute confusion (ex. ne pas presenter la congestion comme "la congestion de YAS").
"""
from __future__ import annotations

import argparse
import datetime as dt
from typing import Optional

from src.output.heatmap import Heatmap, COULEUR_NIVEAU_1, COULEUR_NIVEAU_2, COULEUR_NORMAL


class Rapport:
    """Genere le rapport Excel pour un operateur et un jour donnes, et journalise l'execution."""

    def generer(self, jour: dt.date, operateur_id: str, chemin: str) -> dict:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font
        from src.donnees.db import connexion

        debut = dt.datetime.combine(jour, dt.time.min, tzinfo=dt.timezone.utc)
        fin = debut + dt.timedelta(days=1)

        wb = Workbook()

        # ---------- Feuille 1 : Heatmap (filtree operateur, via forfait.operateur_id) ----------
        with connexion() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT DISTINCT t.forfait_id FROM trafic_horaire t
                    JOIN forfait f ON f.id = t.forfait_id
                    WHERE t.bucket >= %s AND t.bucket < %s AND f.operateur_id = %s
                    ORDER BY t.forfait_id
                    """,
                    (debut, fin, operateur_id),
                )
                forfaits_operateur = [r[0] for r in cur.fetchall()]

                cur.execute(
                    """
                    SELECT e.forfait_id, EXTRACT(HOUR FROM e.bucket)::int, e.niveau
                    FROM evenement_detecte e
                    JOIN forfait f ON f.id = e.forfait_id
                    WHERE e.type_evenement = 'ANOMALIE' AND e.bucket >= %s AND e.bucket < %s
                        AND f.operateur_id = %s
                    """,
                    (debut, fin, operateur_id),
                )
                anomalies_map = {(r[0], r[1]): r[2] for r in cur.fetchall()}

                cur.execute(
                    """
                    SELECT bucket, forfait_id, score, niveau FROM evenement_detecte e
                    JOIN forfait f ON f.id = e.forfait_id
                    WHERE e.type_evenement = 'ANOMALIE' AND e.bucket >= %s AND e.bucket < %s
                        AND f.operateur_id = %s
                    ORDER BY bucket, forfait_id
                    """,
                    (debut, fin, operateur_id),
                )
                lignes_anomalies = cur.fetchall()

                cur.execute(
                    """
                    SELECT bucket, indicateur, score, seuil FROM evenement_detecte
                    WHERE type_evenement = 'ALERTE_QOS' AND bucket >= %s AND bucket < %s
                    ORDER BY bucket, indicateur
                    """,
                    (debut, fin),
                )
                lignes_alertes = cur.fetchall()

                cur.execute(
                    """
                    SELECT bucket, forfait_id, indicateur, score, est_causal FROM evenement_detecte e
                    JOIN forfait f ON f.id = e.forfait_id
                    WHERE e.type_evenement = 'CORRELATION' AND e.bucket >= %s AND e.bucket < %s
                        AND f.operateur_id = %s
                    ORDER BY bucket, forfait_id, indicateur
                    """,
                    (debut, fin, operateur_id),
                )
                lignes_correlations = cur.fetchall()

        # ---------- Sheet 1 : Heatmap ----------
        ws1 = wb.active
        ws1.title = "Heatmap"
        ws1.cell(row=1, column=1, value=f"forfait_id ({operateur_id})")
        for h in range(24):
            ws1.cell(row=1, column=h + 2, value=h)
        fills = {
            2: PatternFill(start_color=COULEUR_NIVEAU_2, end_color=COULEUR_NIVEAU_2, fill_type="solid"),
            1: PatternFill(start_color=COULEUR_NIVEAU_1, end_color=COULEUR_NIVEAU_1, fill_type="solid"),
            None: PatternFill(start_color=COULEUR_NORMAL, end_color=COULEUR_NORMAL, fill_type="solid"),
        }
        for i, fid in enumerate(sorted(forfaits_operateur), start=2):
            ws1.cell(row=i, column=1, value=fid)
            for h in range(24):
                niveau = anomalies_map.get((fid, h))
                cell = ws1.cell(row=i, column=h + 2, value=niveau if niveau is not None else "")
                cell.fill = fills[niveau]

        # ---------- Sheet 2 : Detail anomalies ----------
        ws2 = wb.create_sheet("Anomalies")
        ws2.append(["bucket", "forfait_id", "score (alpha)", "niveau"])
        for row in lignes_anomalies:
            ws2.append([row[0].isoformat(), row[1], float(row[2]), row[3]])

        # ---------- Sheet 3 : Alertes QoS (reseau entier, PAS filtre operateur) ----------
        ws3 = wb.create_sheet("Alertes QoS")
        ws3.cell(row=1, column=1, value="ATTENTION : reseau entier, tous operateurs confondus (QUALCOP ne distingue pas les operateurs)")
        ws3.cell(row=1, column=1).font = Font(bold=True, italic=True)
        ws3.append(["bucket", "indicateur", "score observe", "seuil (3-sigma)"])
        for row in lignes_alertes:
            ws3.append([row[0].isoformat(), row[1], float(row[2]), float(row[3]) if row[3] is not None else None])

        # ---------- Sheet 4 : Correlations ----------
        ws4 = wb.create_sheet("Correlations")
        ws4.cell(row=1, column=1, value="Note : correlation calculee sur l'historique disponible - a interpreter avec prudence si peu de jours accumules")
        ws4.cell(row=1, column=1).font = Font(bold=True, italic=True)
        ws4.append(["bucket", "forfait_id", "indicateur", "score (Pearson r)", "cause probable"])
        for row in lignes_correlations:
            ws4.append([row[0].isoformat(), row[1], row[2], float(row[3]), row[4]])

        wb.save(chemin)

        # ---------- Journalisation ----------
        with connexion() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "INSERT INTO rapport (operateur_id, chemin_excel) VALUES (%s, %s)",
                    (operateur_id, chemin),
                )

        return {
            "operateur_id": operateur_id,
            "jour": jour.isoformat(),
            "forfaits": len(forfaits_operateur),
            "anomalies": len(lignes_anomalies),
            "alertes_qos": len(lignes_alertes),
            "correlations": len(lignes_correlations),
            "chemin": chemin,
        }


def _parse_date(valeur: str) -> dt.date:
    return dt.date.fromisoformat(valeur)


def main(argv: Optional[list[str]] = None) -> None:
    p = argparse.ArgumentParser(description="Genere le rapport Excel (heatmap + detail) pour un operateur et un jour.")
    p.add_argument("--jour", required=True, help="Jour a analyser, ISO (ex. 2025-08-02)")
    p.add_argument("--operateur", required=True, help="Code operateur (ex. YAS, MOOV)")
    p.add_argument("--sortie", default=None, help="Chemin du fichier .xlsx (defaut : rapport_<operateur>_<jour>.xlsx)")
    args = p.parse_args(argv)

    jour = _parse_date(args.jour)
    chemin = args.sortie or f"rapport_{args.operateur}_{jour.isoformat()}.xlsx"

    rapport = Rapport()
    resultat = rapport.generer(jour, args.operateur, chemin)
    print(
        f"Rapport genere : {resultat['operateur_id']} / {resultat['jour']} -> {resultat['chemin']} "
        f"({resultat['forfaits']} forfaits, {resultat['anomalies']} anomalies, "
        f"{resultat['alertes_qos']} alertes QoS, {resultat['correlations']} correlations)"
    )


if __name__ == "__main__":
    main()